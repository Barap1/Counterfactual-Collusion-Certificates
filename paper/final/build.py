#!/usr/bin/env python3
"""Build editable DOCX and vector PDF from canonical manuscript.md; no experiment execution.
Use Python with reportlab, python-docx and pypdf. See EDITING_GUIDE.md.
"""
from pathlib import Path
import argparse,json,hashlib,re,html,os,tempfile
ROOT=Path(__file__).resolve().parents[2]; P=ROOT/'paper/final'; FONT=ROOT/'paper/v2/fonts'
def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def blocks(text):
 result=[]
 for raw in re.split(r'\n\s*\n',text.strip()):
  raw=raw.strip()
  if not raw:continue
  if raw=='<!-- PAGEBREAK -->':result.append(('break',''));continue
  if raw.startswith('<!--'):continue
  if raw.startswith('|'):
   rows=[[c.strip() for c in row.strip().strip('|').split('|')] for row in raw.splitlines() if row.startswith('|')];rows=[r for r in rows if not all(re.fullmatch(r'[:\- ]+',c) for c in r)];result.append(('table',rows));continue
  m=re.match(r'!\[([^]]*)\]\(([^)]+)\)',raw)
  if m:result.append(('image',{'alt':m[1],'path':m[2]}));continue
  m=re.match(r'^(#{1,3}) (.*)',raw,re.S)
  if m:result.append(('h'+str(len(m[1])),m[2]));continue
  result.append(('caption' if raw.startswith('**Figure ') or raw.startswith('**Table ') else 'p',raw.replace('\n',' ')))
 return result

def inline(s):
 s=html.escape(s)
 s=re.sub(r'\[([^]]+)\]\(([^)]+)\)',r'<a href="\2" color="#193025">\1</a>',s)
 s=re.sub(r'\*\*(.*?)\*\*',r'<b>\1</b>',s)
 s=re.sub(r'`([^`]+)`',r'<font name="Courier" size="9">\1</font>',s)
 return s

def pdf_build(bs,out):
 from reportlab.platypus import SimpleDocTemplate,BaseDocTemplate,PageTemplate,Frame,Paragraph,Table,TableStyle,PageBreak,Flowable,KeepTogether
 from reportlab.lib.styles import ParagraphStyle
 from reportlab.pdfbase import pdfmetrics
 from reportlab.pdfbase.ttfonts import TTFont
 from reportlab.lib import colors
 from pypdf import PdfReader,PdfWriter,Transformation
 for name,f in [('OST','OldStandard-Regular.ttf'),('OSTB','OldStandard-Bold.ttf'),('OSTI','OldStandard-Italic.ttf')]:pdfmetrics.registerFont(TTFont(name,str(FONT/f)))
 pdfmetrics.registerFontFamily('OST',normal='OST',bold='OSTB',italic='OSTI',boldItalic='OSTB')
 styles={
 'p':ParagraphStyle('p',fontName='OST',fontSize=11,leading=12.65,spaceAfter=4,textColor=colors.black),
 'h1':ParagraphStyle('title',fontName='OSTB',fontSize=20,leading=23,alignment=1,spaceAfter=12),
 'h2':ParagraphStyle('h2',fontName='OSTB',fontSize=14,leading=21,spaceBefore=10,spaceAfter=4,keepWithNext=True),
 'h3':ParagraphStyle('h3',fontName='OSTB',fontSize=13,leading=16,spaceBefore=8,spaceAfter=4,keepWithNext=True),
 'caption':ParagraphStyle('cap',fontName='OST',fontSize=10,leading=11.5,spaceAfter=8),
 'cell':ParagraphStyle('cell',fontName='OST',fontSize=10,leading=11.5),
 'abstract':ParagraphStyle('abstract',fontName='OSTI',fontSize=11,leading=12.65,leftIndent=27,rightIndent=27,spaceAfter=10),
 'author':ParagraphStyle('author',fontName='OST',fontSize=11,leading=13,alignment=1,spaceAfter=4)}
 overlays=[];story=[];in_abstract=False;before_abstract=True
 class Fig(Flowable):
  def __init__(self,val):
   super().__init__();self.path=P/val['path'];svg=self.path.read_text();m=re.search(r'viewBox="0 0 ([\d.]+) ([\d.]+)"',svg);self.width=468;self.height=468*float(m[2])/float(m[1]);self.spaceAfter=6
  def draw(self):
   x,y=self.canv.absolutePosition(0,0);overlays.append({'page':self.canv.getPageNumber()-1,'x':x,'y':y,'width':self.width,'height':self.height,'path':str(self.path.with_suffix('.pdf'))})
 def footer(c,doc):
  c.setFont('OST',9);c.drawCentredString(306,36,str(doc.page));
  if doc.page==1:c.setStrokeColor(colors.black);c.setLineWidth(2);c.line(72,720,540,720)
 for typ,val in bs:
  if typ=='break':story.append(PageBreak());continue
  if typ=='image':story.append(Fig(val));continue
  if typ=='table':
   data=[[Paragraph(inline(c),styles['cell']) for c in row] for row in val];widths=[90,162,216] if len(val[0])==3 else [91,91,101,113,72]
   t=Table(data,colWidths=widths,hAlign='LEFT',repeatRows=1);t.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#f2f5f2')),('LINEABOVE',(0,0),(-1,0),.7,colors.black),('LINEBELOW',(0,0),(-1,0),.5,colors.grey),('LINEBELOW',(0,-1),(-1,-1),.7,colors.black),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]));story.append(t);continue
  style=styles.get(typ,styles['p'])
  if typ=='h2' and val=='Abstract':
   style=ParagraphStyle('abshead',parent=styles['h2'],alignment=1,spaceBefore=8);in_abstract=True;before_abstract=False
  elif typ in ['h2','h3']:in_abstract=False
  elif in_abstract and typ=='p':style=styles['abstract']
  elif before_abstract and typ=='p':style=styles['author'];val=val.replace('Georgia Institute of Technology','<br/>Georgia Institute of Technology').replace('With Apart Research','<br/>With Apart Research').replace('AI Collusion Research Sprint','<br/>AI Collusion Research Sprint')
  story.append(Paragraph(inline(val).replace('&lt;br/&gt;','<br/>'),style))
 temp=P/'qa/layout-unmerged.pdf';doc=BaseDocTemplate(str(temp),pagesize=(612,792),leftMargin=72,rightMargin=72,topMargin=72,bottomMargin=72,title='Same Evidence, Different Verdicts',author='Aarav Patel',pageCompression=1);doc.addPageTemplates(PageTemplate(id='body',frames=[Frame(72,72,468,648,leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)],onPage=footer));doc.build(story)
 reader=PdfReader(temp);writer=PdfWriter()
 for i,page in enumerate(reader.pages):
  for ov in overlays:
   if ov['page']==i:
    fig=PdfReader(ov['path']).pages[0];sx=ov['width']/float(fig.mediabox.width);sy=ov['height']/float(fig.mediabox.height);page.merge_transformed_page(fig,Transformation().scale(sx,sy).translate(ov['x'],ov['y']))
  writer.add_page(page)
 writer.add_metadata({'/Title':'Same Evidence, Different Verdicts: Evidence Standards for AI Collusion Attribution','/Author':'Aarav Patel'})
 with open(out,'wb') as f:writer.write(f)
 return {'pages':len(reader.pages),'figure_placements':overlays}

def docx_build(bs,out):
 from docx import Document
 from docx.shared import Inches,Pt,RGBColor
 from docx.enum.text import WD_ALIGN_PARAGRAPH
 from docx.enum.style import WD_STYLE_TYPE
 from docx.oxml import OxmlElement
 from docx.oxml.ns import qn
 # Retain native style definitions/section geometry from the inspected official export.
 template=Path('/tmp/ccc-official-template.docx');doc=Document(str(template)) if template.exists() else Document()
 body=doc._element.body
 for child in list(body):
  if child.tag!=qn('w:sectPr'):body.remove(child)
 sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1);sec.header_distance=sec.footer_distance=Inches(.5)
 for name,size in [('Normal',11),('Title',20),('Heading 1',14),('Heading 2',13),('Caption',10)]:
  st=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH);st.font.name='Old Standard TT';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.line_spacing=1.15;st.paragraph_format.space_after=Pt(4)
  if name!='Normal':st.font.bold=name!='Caption'
 doc.styles['Heading 1'].paragraph_format.line_spacing=1.5;doc.styles['Heading 1'].paragraph_format.space_before=Pt(10)
 doc.styles['Heading 2'].paragraph_format.space_before=Pt(8)
 def runs(p,text):
  # Native editable emphasis and hyperlinks (not flattened page images).
  for part in re.split(r'(\*\*.*?\*\*|`[^`]+`|\[[^]]+\]\([^)]+\))',text):
   m=re.fullmatch(r'\[([^]]+)\]\(([^)]+)\)',part)
   if m:
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(m[2],'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True));r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=m[1];r.append(t);h.append(r);p._p.append(h)
   else:
    r=p.add_run(part.strip('*`'));r.bold=part.startswith('**');
    if part.startswith('`'):r.font.name='Courier New';r.font.size=Pt(9)
 in_abs=False;before_abs=True
 for typ,val in bs:
  if typ=='break':doc.add_page_break();continue
  if typ=='image':
   p=doc.add_paragraph();p.paragraph_format.space_after=Pt(6);shape=p.add_run().add_picture(str((P/val['path']).with_suffix('.png')),width=Inches(6.5));shape._inline.docPr.set('descr',val['alt']);continue
  if typ=='table':
   t=doc.add_table(rows=0,cols=len(val[0]));t.autofit=False;widths=[90,162,216] if len(val[0])==3 else [91,91,101,113,72]
   for rownum,row in enumerate(val):
    cells=t.add_row().cells
    for j,c in enumerate(row):
     cells[j].width=Pt(widths[j]);p=cells[j].paragraphs[0];runs(p,c)
     for r in p.runs:r.font.size=Pt(10);r.bold=rownum==0
     p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.15
    if rownum==0:
     pr=t.rows[0]._tr.get_or_add_trPr();pr.append(OxmlElement('w:tblHeader'))
   continue
  style={'h1':'Title','h2':'Heading 1','h3':'Heading 2','caption':'Caption'}.get(typ,'Normal');p=doc.add_paragraph(style=style);runs(p,val)
  if typ=='h1':p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  if typ=='h2' and val=='Abstract':p.alignment=WD_ALIGN_PARAGRAPH.CENTER;in_abs=True;before_abs=False
  elif typ in ['h2','h3']:in_abs=False
  elif in_abs:
   p.paragraph_format.left_indent=Pt(27);p.paragraph_format.right_indent=Pt(27)
   for r in p.runs:r.italic=True
  elif before_abs and typ=='p':p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 doc.core_properties.author='Aarav Patel';doc.core_properties.title='Same Evidence, Different Verdicts';doc.core_properties.subject='Counterfactual evidence for collusion attribution';doc.save(out)

def main():
 ap=argparse.ArgumentParser();ap.add_argument('--check',action='store_true');args=ap.parse_args();inputs=[P/'manuscript.md',Path(__file__),*FONT.glob('*.ttf'),*sorted((P/'figures').glob('fig*.svg')),*sorted((P/'figures').glob('fig*.pdf')),*sorted((P/'figures').glob('fig*.png'))];hashes={str(p.relative_to(ROOT)):digest(p) for p in inputs};receipt=P/'qa/build_manifest.json'
 if args.check:
  old=json.loads(receipt.read_text());ok=old['inputs']==hashes and all((ROOT/p).exists() and digest(ROOT/p)==h for p,h in old['outputs'].items());print(json.dumps({'current':ok,'pages':old['pdf']['pages']}));raise SystemExit(0 if ok else 1)
 bs=blocks((P/'manuscript.md').read_text());(P/'build').mkdir(exist_ok=True);pdf=P/'build/local_submission.pdf';dx=P/'build/local_submission.docx'
 info=pdf_build(bs,pdf);docx_build(bs,dx);receipt.write_text(json.dumps({'inputs':hashes,'outputs':{str(p.relative_to(ROOT)):digest(p) for p in [pdf,dx]},'pdf':info,'source':'paper/final/manuscript.md','format':'local template adaptation; native Google Doc separately exported/verified'},indent=2)+'\n');print(json.dumps(info))
if __name__=='__main__':main()
